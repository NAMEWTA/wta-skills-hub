#!/usr/bin/env python3
"""Render a tailored Markdown resume to HTML, and DOCX/PDF when libraries exist."""

from __future__ import annotations

import argparse
import html
import json
import re
import sys
from pathlib import Path

CLAIM_COMMENT = re.compile(r"<!--\s*claim:C-\d+\s*-->", re.IGNORECASE)


def md_to_blocks(text: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(("p", " ".join(paragraph).strip()))
            paragraph.clear()

    for raw in text.splitlines():
        line = CLAIM_COMMENT.sub("", raw).rstrip()
        if not line.strip() or line.strip().startswith("<!--"):
            flush()
            continue
        if line.startswith("### "):
            flush()
            blocks.append(("h3", line[4:].strip()))
        elif line.startswith("## "):
            flush()
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("# "):
            flush()
            blocks.append(("h1", line[2:].strip()))
        elif line.startswith("- "):
            flush()
            blocks.append(("li", line[2:].strip()))
        else:
            paragraph.append(line.strip())
    flush()
    return blocks


def to_html(blocks: list[tuple[str, str]], css: str) -> str:
    parts = [
        "<!DOCTYPE html>",
        '<html lang="zh">',
        "<head>",
        '<meta charset="utf-8">',
        "<title>Resume</title>",
        f"<style>\n{css}\n</style>",
        "</head>",
        "<body>",
        '<main class="sheet">',
    ]
    list_open = False
    for kind, content in blocks:
        safe = html.escape(content)
        if kind != "li" and list_open:
            parts.append("</ul>")
            list_open = False
        if kind == "li":
            if not list_open:
                parts.append("<ul>")
                list_open = True
            parts.append(f"<li>{safe}</li>")
        elif kind in {"h1", "h2", "h3", "p"}:
            parts.append(f"<{kind}>{safe}</{kind}>")
    if list_open:
        parts.append("</ul>")
    parts.extend(["</main>", "</body>", "</html>", ""])
    return "\n".join(parts)


def try_docx(blocks: list[tuple[str, str]], dest: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        return "skipped: python-docx not installed"
    doc = Document()
    for kind, content in blocks:
        if kind == "h1":
            doc.add_heading(content, level=0)
        elif kind == "h2":
            doc.add_heading(content, level=1)
        elif kind == "h3":
            doc.add_heading(content, level=2)
        elif kind == "li":
            doc.add_paragraph(content, style="List Bullet")
        else:
            doc.add_paragraph(content)
    doc.save(dest)
    return "written"


def try_pdf(html_path: Path, dest: Path) -> str:
    try:
        from weasyprint import HTML
    except ImportError:
        return "skipped: weasyprint not installed"
    HTML(filename=str(html_path)).write_pdf(str(dest))
    return "written"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--md", required=True)
    parser.add_argument("--css", required=True)
    parser.add_argument("--out-dir", required=True)
    args = parser.parse_args()

    md_path = Path(args.md)
    css_path = Path(args.css)
    out_dir = Path(args.out_dir)
    if not md_path.is_file():
        print(f"missing markdown: {md_path}", file=sys.stderr)
        raise SystemExit(2)
    css = css_path.read_text(encoding="utf-8") if css_path.is_file() else "body{font-family:sans-serif}"
    out_dir.mkdir(parents=True, exist_ok=True)
    blocks = md_to_blocks(md_path.read_text(encoding="utf-8"))
    html_path = out_dir / "resume-tailored.html"
    html_path.write_text(to_html(blocks, css), encoding="utf-8")
    docx_status = try_docx(blocks, out_dir / "resume-tailored.docx")
    pdf_status = try_pdf(html_path, out_dir / "resume-tailored.pdf")
    status = {
        "html": "written",
        "docx": docx_status,
        "pdf": pdf_status,
    }
    (out_dir / "render-status.json").write_text(
        json.dumps(status, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(status, ensure_ascii=False))
    if docx_status != "written" and pdf_status != "written":
        raise SystemExit(3)


if __name__ == "__main__":
    main()
